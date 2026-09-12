from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from orion.models import DocumentChunk, SourceType

def parse_pdf(path: Path, document_id: str) -> list[DocumentChunk]:
    reader = PdfReader(path)

    chunks: list[DocumentChunk] = []

    for page_number, page in enumerate(reader.pages, start = 1):
        text = page.extract_text() or ""

        if not text.strip():
            continue

        chunks.append(
            DocumentChunk(
                document_id = document_id,
                document_name = path.name,
                source_type = SourceType.PDF,
                page = page_number,
                text = text.strip()
            )
        )

    return chunks

def parse_docx(path: Path, document_id: str) -> list[DocumentChunk]:
    document = Document(path)

    chunks: list[DocumentChunk] = []

    for paragraph_number, paragraph in enumerate(document.paragraphs, start = 1):
        text = paragraph.text.strip()

        if not text:
            continue

        chunks.append(
            DocumentChunk(
                document_id = document_id,
                document_name = path.name,
                source_type = SourceType.DOCX,
                section = f"paragraph-{paragraph_number}",
                text = text
            )
        )

    return chunks

def parse_xlsx(path: Path, document_id: str) -> list[DocumentChunk]:
    workbook = load_workbook(
        path,
        read_only = True,
        data_only = True # read stored calculated text instead of formula text
    )

    chunks: list[DocumentChunk] = []

    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only = True), start = 1):
            values = [str(value).strip() for value in row if value is not None]

            if not values:
                continue

            text = " | ".join(values)

            chunks.append(
                DocumentChunk(
                    document_id = document_id,
                    document_name = path.name,
                    source_type = SourceType.XLSX,
                    section = f"{sheet.title}:row-{row_number}",
                    text = text
                )
            )

    return chunks

def parse_document(path: Path, document_id: str) -> list[DocumentChunk]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return parse_pdf(path, document_id)

    if suffix == ".docx":
        return parse_docx(path, document_id)

    if suffix == ".xlsx":
        return parse_xlsx(path, document_id)

    raise ValueError(
        f"Unsupported document type: {suffix}"
    )